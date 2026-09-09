"""
================================================================================
document_loader.py  —— 数据加载 & 分割（一个文件搞定，不依赖 RAG-Anything 库）
================================================================================
两段流水线：
  ① 加载/解析：PDF / 图片 用【MinerU 本地模型】解析；Word(.docx) 用 python-docx；
              txt/md 直接读。全部统一输出成同一种中间格式 content_list。
  ② 分割/分块：把 content_list 做“多模态分块”——正文按章节/长度切块，
              图片/表格/公式各自成块并自动带上所属章节与前后正文。

多模态分块算法思路抽自 RAG-Anything（utils.separate_content、
modalprocessors.ContextExtractor、processor 的成块模板），这里是去框架依赖的独立实现。

--------------------------------------------------------------------------------
【依赖】按格式安装（纯本地、不需要任何 API KEY）：
  PDF/图片（MinerU，本地免费）：pip install -U "mineru[core]"     # 首次运行会自动下载模型
  Word：                        pip install python-docx
  txt/md、分块本身：            无需第三方库
注：tiktoken 可选，装了按真实 token 计块长，不装按字符数，不影响运行。

【最快验证】python document_loader.py    （离线，无需 MinerU/Key，跑通解析→分块逻辑）
【真实使用】见文件末尾 Example，或直接：
    loader = DocumentLoader()
    chunks = loader.load_and_split("某文件.pdf")
================================================================================
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any, Awaitable, Callable, Dict, List, Optional, Tuple

# ###########################################################################
# 0. 通用小工具
# ###########################################################################
try:
    import tiktoken
    _ENC = tiktoken.get_encoding("cl100k_base")

    def count_tokens(text: str) -> int:
        return len(_ENC.encode(text or ""))
except Exception:
    def count_tokens(text: str) -> int:   # 没装 tiktoken 就按字符数
        return len(text or "")


def chunk_id_of(text: str) -> str:
    return "chunk-" + hashlib.md5(text.encode("utf-8")).hexdigest()[:16]


# ###########################################################################
# ============================ ① 数据加载 / 统一解析 =========================
# ###########################################################################

# ----------------------------- txt / md（零依赖） -----------------------------
class TextParser:
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        path = Path(file_path)
        text = path.read_text(encoding="utf-8", errors="ignore")
        out: List[Dict[str, Any]] = []
        if path.suffix.lower() == ".md":
            for line in text.splitlines():
                s = line.strip()
                if not s:
                    continue
                m = re.match(r"^(#{1,6})\s+(.*)$", s)
                if m:
                    out.append({"type": "text", "text": m.group(2).strip(),
                                "text_level": len(m.group(1)), "page_idx": 0})
                else:
                    out.append({"type": "text", "text": s, "text_level": 0, "page_idx": 0})
        else:
            for para in re.split(r"\n\s*\n", text):
                p = para.strip()
                if p:
                    out.append({"type": "text", "text": p, "text_level": 0, "page_idx": 0})
        return out


# ----------------------------- Word .docx（python-docx） -----------------------------
class DocxParser:
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        try:
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as e:
            raise ImportError("解析 Word 请先安装：pip install python-docx") from e

        doc = Document(str(file_path))
        out: List[Dict[str, Any]] = []
        for child in doc.element.body.iterchildren():  # 按原文顺序读段落与表格
            if child.tag == qn("w:p"):
                p = Paragraph(child, doc)
                text = p.text.strip()
                if text:
                    style = p.style.name if p.style is not None else ""
                    out.append({"type": "text", "text": text,
                                "text_level": self._heading_level(style), "page_idx": 0})
            elif child.tag == qn("w:tbl"):
                md = self._table_to_markdown(Table(child, doc))
                if md.strip():
                    out.append({"type": "table", "table_body": md, "page_idx": 0})
        return out

    @staticmethod
    def _heading_level(style_name: str) -> int:
        m = re.search(r"(?:heading|标题)\s*(\d)", style_name or "", re.IGNORECASE)
        return int(m.group(1)) if m else 0

    @staticmethod
    def _table_to_markdown(table) -> str:
        rows = [[c.text.strip().replace("\n", " ").replace("|", "/") for c in r.cells]
                for r in table.rows]
        if not rows:
            return ""
        ncol = max(len(r) for r in rows)
        head = rows[0] + [""] * (ncol - len(rows[0]))
        lines = ["| " + " | ".join(head) + " |",
                 "| " + " | ".join(["---"] * ncol) + " |"]
        for r in rows[1:]:
            r = r + [""] * (ncol - len(r))
            lines.append("| " + " | ".join(r) + " |")
        return "\n".join(lines)


# ----------------------------- PDF / 图片：MinerU 本地解析 -----------------------------
class MinerUParser:
    """调用本地 MinerU 命令行解析 PDF/图片，读取其 content_list.json 并归一化。

    MinerU 是开源免费的本地模型（不是收费云 API）：
      安装：pip install -U "mineru[core]"（首次解析会自动下载模型权重）
      无 GPU 用 backend="pipeline"（纯 CPU）；有 GPU 可用默认/hybrid 更快。
    """

    # MinerU 输出里属于“页面噪声”的类型，直接丢弃，不进知识库
    NOISE_TYPES = {"header", "footer", "page_number", "aside_text", "page_footnote"}

    def __init__(self, output_dir: str = "./mineru_out", backend: str = "pipeline",
                 method: str = "auto", lang: str = "ch"):
        self.output_dir = Path(output_dir)
        self.backend = backend      # pipeline(CPU兼容) / vlm-transformers / hybrid-engine
        self.method = method        # auto / txt / ocr（仅 pipeline 生效）
        self.lang = lang

    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(file_path)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        # 1) 调 MinerU 命令行（等价于手动执行 mineru -p xx -o xx -m auto -b pipeline）
        cmd = ["mineru", "-p", str(path), "-o", str(self.output_dir),
               "-m", self.method, "-b", self.backend, "-l", self.lang]
        try:
            subprocess.run(cmd, check=True)
        except FileNotFoundError as e:
            raise RuntimeError(
                "未找到 mineru 命令，请先安装本地解析模型：pip install -U \"mineru[core]\""
            ) from e
        except subprocess.CalledProcessError as e:
            raise RuntimeError(f"MinerU 解析失败：{e}") from e

        # 2) 找到输出的 {stem}_content_list.json（不同后端目录层级略有差异，用 rglob 兼容）
        matches = sorted(self.output_dir.rglob(f"{path.stem}_content_list.json"))
        if not matches:
            matches = sorted(self.output_dir.rglob("*content_list.json"))
        if not matches:
            raise RuntimeError("MinerU 已运行但未找到 content_list.json，请检查输出目录。")
        cl_path = matches[-1]
        raw = json.loads(cl_path.read_text(encoding="utf-8"))

        # 3) 归一化成我们统一的 content_list
        return self.normalize(raw, cl_path.parent)

    @staticmethod
    def _clean_table_body(body: str) -> str:
        # MinerU 表格可能裹着 <html><body> 外壳，剥掉，保留 table 本身
        return re.sub(r"(?is)</?(?:html|body)[^>]*>", "", body or "").strip()

    def normalize(self, raw: List[Dict[str, Any]], base_dir: Path) -> List[Dict[str, Any]]:
        """把 MinerU 的 content_list 归一化：去噪声、图片相对路径转绝对、特殊类型并入统一格式。"""
        out: List[Dict[str, Any]] = []
        for it in raw:
            if not isinstance(it, dict):
                continue
            t = it.get("type", "text")
            if t in self.NOISE_TYPES:                 # 页眉/页脚/页码等噪声丢弃
                continue

            if t in ("image", "chart", "table", "equation"):
                ip = it.get("img_path")
                if ip:                                # MinerU 给的是 images/xx.jpg 相对路径
                    it["img_path"] = str((base_dir / ip).resolve())
                if t == "chart":                      # 统计图表按图片处理
                    it["type"] = "image"
                if t == "table":
                    it["table_body"] = self._clean_table_body(it.get("table_body", ""))
                out.append(it)
            elif t == "text":
                if (it.get("text") or "").strip():
                    out.append(it)
            elif t == "code":                         # 代码块并入正文
                body = it.get("code_body", "")
                if body.strip():
                    out.append({"type": "text", "text": body, "page_idx": it.get("page_idx", 0)})
            elif t == "list":                         # 列表项拼成正文
                items = it.get("list_items", [])
                if items:
                    out.append({"type": "text", "text": "\n".join(items),
                                "page_idx": it.get("page_idx", 0)})
            else:                                     # 其它带 text 的兜底保留
                tx = it.get("text")
                if isinstance(tx, str) and tx.strip():
                    it["type"] = "text"
                    out.append(it)
        return out


# ----------------------------- 统一解析入口：任何格式 → content_list -----------------------------
class UnifiedDocumentParser:
    IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp", ".gif"}
    TEXT_EXTS = {".txt", ".md"}
    DOCX_EXTS = {".docx"}
    MINERU_EXTS = {".pdf"} | IMAGE_EXTS

    def __init__(self, mineru_output_dir: str = "./mineru_out", backend: str = "pipeline",
                 method: str = "auto", lang: str = "ch"):
        self.text_parser = TextParser()
        self.docx_parser = DocxParser()
        self.mineru_parser = MinerUParser(mineru_output_dir, backend, method, lang)

    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        ext = Path(file_path).suffix.lower()
        if ext in self.TEXT_EXTS:
            return self.text_parser.parse(file_path)
        if ext in self.DOCX_EXTS:
            return self.docx_parser.parse(file_path)
        if ext in self.MINERU_EXTS:
            return self.mineru_parser.parse(file_path)
        raise NotImplementedError(
            f"暂不支持 {ext}。.doc/.ppt/.xls 请先用 Office 另存为 .pdf/.docx；"
            f"当前支持 pdf/图片/docx/txt/md。")


# ###########################################################################
# ============================ ② 多模态分块 =========================
# ###########################################################################
def extract_section_path(content_list: List[Dict[str, Any]], current_index: int) -> str:
    """向前扫描标题，用栈还原 章 > 节 路径（对应 RAG-Anything 同名函数）。"""
    if not content_list or current_index is None:
        return ""
    chain: List[Tuple[int, str]] = []
    for item in content_list[:max(0, int(current_index))]:
        if item.get("type", "text") != "text":
            continue
        text = str(item.get("text", "") or "").strip()
        if not text:
            continue
        try:
            level = int(item.get("text_level", 0) or 0)
        except (TypeError, ValueError):
            level = 0
        if level <= 0:
            continue
        while chain and chain[-1][0] >= level:
            chain.pop()
        chain.append((level, text))
    return " > ".join(t for _, t in chain)


def extract_neighbor_text(content_list: List[Dict[str, Any]], idx: int, window: int = 3) -> str:
    if not content_list:
        return ""
    parts = []
    for pos in range(max(0, idx - window), min(len(content_list), idx + window + 1)):
        if pos == idx:
            continue
        item = content_list[pos]
        if item.get("type", "text") == "text" and str(item.get("text", "")).strip():
            parts.append(str(item["text"]).strip())
    return " ".join(parts)


def separate_content(content_list):
    """拆成【文本段序列】和【多模态项】，多模态项预计算章节路径与相邻正文。"""
    text_segments, mm_items = [], []
    for index, item in enumerate(content_list):
        if item.get("type", "text") == "text":
            text = str(item.get("text", "") or "")
            if text.strip():
                text_segments.append({"text": text,
                                      "level": int(item.get("text_level", 0) or 0),
                                      "page_idx": item.get("page_idx", 0), "index": index})
        else:
            mm = dict(item)
            mm["_content_list_index"] = index
            mm.setdefault("_section_path", extract_section_path(content_list, index))
            mm.setdefault("_neighbor_text", extract_neighbor_text(content_list, index))
            mm_items.append(mm)
    return text_segments, mm_items


_BREAKS = ["\n\n", "\n", "。", "！", "？", "；", ". ", "! ", "? ", "; ", " "]


def _hard_split(text: str, size: int, overlap: int = 0) -> List[str]:
    pieces, start, n = [], 0, len(text)
    while start < n:
        end = min(n, start + size)
        window = text[start:end]
        if end < n:
            for br in _BREAKS:
                pos = window.rfind(br, int(size * 0.6))
                if pos != -1:
                    end = start + pos + len(br)
                    window = text[start:end]
                    break
        if window.strip():
            pieces.append(window.strip())
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return pieces


def chunk_segments(segments, chunk_size=1200, overlap=120):
    """文本按“尽量不跨标题 + 长度上限 + 块间重叠”打包。"""
    chunks, buf, buf_len, buf_idx, buf_page = [], [], 0, [], []

    def flush():
        nonlocal buf, buf_len, buf_idx, buf_page
        if not buf:
            return
        joined = "\n\n".join(buf)
        chunks.append({"text": joined, "start_index": min(buf_idx),
                       "page_idx": min(buf_page) if buf_page else 0})
        if overlap and joined:
            tail = joined[-overlap:]
            buf, buf_len = [tail], count_tokens(tail)
            p = min(buf_page) if buf_page else 0
            buf_idx, buf_page = [min(buf_idx)], [p]
        else:
            buf, buf_len, buf_idx, buf_page = [], 0, [], []

    for seg in segments:
        level = seg["level"]
        rendered = ("#" * level + " " + seg["text"]) if level > 0 else seg["text"]
        seg_len = count_tokens(rendered)
        if buf and level > 0 and buf_len + seg_len > chunk_size and buf_len >= chunk_size * 0.5:
            flush()
        if seg_len > chunk_size:
            flush()
            for piece in _hard_split(rendered, chunk_size, overlap):
                chunks.append({"text": piece, "start_index": seg["index"], "page_idx": seg["page_idx"]})
            continue
        if buf_len + seg_len > chunk_size:
            flush()
        buf.append(rendered)
        buf_len += seg_len
        buf_idx.append(seg["index"])
        buf_page.append(int(seg.get("page_idx", 0) or 0))
    flush()
    return chunks


class ContextWindow:
    """给图/表/公式取前后文（对应 ContextExtractor）。"""
    def __init__(self, window=1, mode="page", max_tokens=1200, include_headers=True):
        self.window, self.mode, self.max_tokens = window, mode, max_tokens
        self.include_headers = include_headers

    def _item_text(self, item):
        if item.get("type", "text") != "text":
            return ""
        text = str(item.get("text", "") or "")
        level = int(item.get("text_level", 0) or 0)
        return ("#" * level + " " + text) if self.include_headers and level > 0 else text

    def extract(self, content_list, item):
        texts = []
        if self.mode == "page":
            cur = int(item.get("page_idx", 0) or 0)
            for it in content_list:
                p = int(it.get("page_idx", 0) or 0)
                if cur - self.window <= p < cur + self.window + 1 and it.get("type", "text") == "text":
                    t = self._item_text(it)
                    if t.strip():
                        texts.append(f"[Page {p}] {t}" if p != cur else t)
        else:
            cur = int(item.get("_content_list_index", 0) or 0)
            for i in range(max(0, cur - self.window), min(len(content_list), cur + self.window + 1)):
                if i == cur:
                    continue
                t = self._item_text(content_list[i])
                if t.strip():
                    texts.append(t)
        return self._truncate("\n".join(texts))

    def _truncate(self, text):
        if not text or len(text) <= self.max_tokens:
            return text
        cut = text[:self.max_tokens]
        for br in ("。", ". ", "\n"):
            p = cut.rfind(br, int(self.max_tokens * 0.8))
            if p != -1:
                return cut[:p + len(br)]
        return cut + " …"


def _captions(item, *keys):
    for k in keys:
        v = item.get(k)
        if isinstance(v, list) and v:
            return ", ".join(str(x) for x in v)
        if isinstance(v, str) and v.strip():
            return v
    return "无"


def render_modal_chunk(item, caption=""):
    """把图/表/公式组装成一个独立块文本（对应 _apply_chunk_template）。"""
    ctype = item.get("type", "generic")
    section = item.get("_section_path", "") or "无"
    neighbor = item.get("_neighbor_text", "") or "无"
    if ctype == "image":
        body = (f"【图片】\n所属章节：{section}\n相邻正文：{neighbor}\n"
                f"图片路径：{item.get('img_path', '')}\n"
                f"图注：{_captions(item, 'image_caption', 'img_caption')}\n"
                f"图片描述：{caption or '（MinerU 已裁剪图片，未生成额外语义描述）'}")
    elif ctype == "table":
        body = (f"【表格】\n所属章节：{section}\n相邻正文：{neighbor}\n"
                f"表注：{_captions(item, 'table_caption')}\n"
                f"表格内容：\n{item.get('table_body', '')}\n"
                f"表格解读：{caption or '（MinerU 已结构化，未额外生成解读）'}")
    elif ctype == "equation":
        eq = item.get("text") or item.get("latex") or ""
        body = (f"【公式】\n所属章节：{section}\n相邻正文：{neighbor}\n"
                f"公式：{eq}\n公式说明：{caption or '（MinerU 已识别为 LaTeX）'}")
    else:
        body = f"【{ctype}】\n所属章节：{section}\n内容：{item}\n描述：{caption}"
    return body, ctype


class MultimodalChunker:
    def __init__(self, chunk_size=1200, overlap=120, context_window=1,
                 context_mode="page", max_context_tokens=1200, caption_func=None):
        self.chunk_size, self.overlap = chunk_size, overlap
        self.context = ContextWindow(context_window, context_mode, max_context_tokens)
        self.caption_func = caption_func   # async (item, ctype, context)->str

    async def chunk(self, content_list: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        segments, mm_items = separate_content(content_list)
        results = [{
            "chunk_id": chunk_id_of(c["text"]), "order_sort": c["start_index"],
            "type": "text", "text": c["text"], "page_idx": c["page_idx"],
            "section_path": extract_section_path(content_list, c["start_index"]),
            "metadata": {},
        } for c in chunk_segments(segments, self.chunk_size, self.overlap)]

        for item in mm_items:
            ctx = self.context.extract(content_list, item)
            caption = ""
            if self.caption_func is not None:
                ret = self.caption_func(item, item.get("type", "generic"), ctx)
                if asyncio.iscoroutine(ret):
                    ret = await ret
                caption = str(ret or "")
            body, ctype = render_modal_chunk(item, caption)
            results.append({
                "chunk_id": chunk_id_of(body), "order_sort": item["_content_list_index"],
                "type": ctype, "text": body, "page_idx": int(item.get("page_idx", 0) or 0),
                "section_path": item.get("_section_path", ""),
                "metadata": {k: v for k, v in item.items() if not k.startswith("_") and k != "type"},
            })

        results.sort(key=lambda x: x["order_sort"])
        for i, r in enumerate(results):
            r["order_index"] = i
            r.pop("order_sort", None)
        return results


# ###########################################################################
# ============================ 门面：加载 + 分割一条龙 =========================
# ###########################################################################
class DocumentLoader:
    def __init__(self, chunk_size: int = 1200, overlap: int = 120, context_window: int = 1,
                 mineru_output_dir: str = "./mineru_out",
                 backend: str = "pipeline", method: str = "auto", lang: str = "ch"):
        """纯本地加载与分割，不需要任何 API KEY。
        backend：MinerU 后端，无 GPU 用 "pipeline"（纯 CPU），有 GPU 可改 "hybrid-engine"。
        """
        self.parser = UnifiedDocumentParser(mineru_output_dir, backend, method, lang)
        self.chunker = MultimodalChunker(chunk_size, overlap, context_window,
                                         caption_func=None)

    # 只解析成统一 content_list
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        return self.parser.parse(file_path)

    # 只对已有 content_list 分块
    async def asplit(self, content_list: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        return await self.chunker.chunk(content_list)

    # 解析 + 分块（异步，FastAPI 里用这个）
    async def aload_and_split(self, file_path: str) -> List[Dict[str, Any]]:
        content_list = await asyncio.to_thread(self.parser.parse, file_path)
        return await self.chunker.chunk(content_list)

    # 解析 + 分块（同步，普通脚本/Jupyter 外用；已有事件循环环境请改用 aload_and_split）
    def load_and_split(self, file_path: str) -> List[Dict[str, Any]]:
        try:
            loop = asyncio.get_running_loop()
        except RuntimeError:
            loop = None
        if loop and loop.is_running():
            raise RuntimeError("检测到已有事件循环（FastAPI/Jupyter），请改用 await loader.aload_and_split(path)")
        return asyncio.run(self.aload_and_split(file_path))


# ###########################################################################
# 离线自测：python document_loader.py
# 不装 MinerU/不联网也能验证：txt/md/docx 解析、以及“MinerU 输出的归一化 + 分块”
# ###########################################################################
def _demo():
    import tempfile

    tmp = Path(tempfile.mkdtemp(prefix="doc_loader_"))
    loader = DocumentLoader(chunk_size=600, overlap=80)  # 不开千问，纯离线

    # 1) markdown 解析→分块
    md = tmp / "a.md"
    md.write_text("# 仓储管理规范\n\n## 入库流程\n\n货物到仓先验收再上架。\n\n"
                  "## 出库流程\n\n先进先出拣货，复核后发运。\n", encoding="utf-8")
    cl = loader.parse(str(md))
    chunks = asyncio.run(loader.asplit(cl))
    print("=" * 70, f"\n[md] 解析出 {len(cl)} 项 → 分出 {len(chunks)} 块")
    for c in chunks:
        print(f"  #{c['order_index']} {c['type']}: {c['text'][:80].replace(chr(10), ' ')}")

    # 2) docx 解析→分块（装了 python-docx 才会跑）
    try:
        from docx import Document
        d = Document()
        d.add_heading("第3章 仓储管理", level=1)
        d.add_heading("3.1 入库", level=2)
        d.add_paragraph("验收后上架。")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "环节", "时限"
        t.cell(1, 0).text, t.cell(1, 1).text = "验收", "2小时"
        dp = tmp / "b.docx"
        d.save(str(dp))
        cl = loader.parse(str(dp))
        chunks = asyncio.run(loader.asplit(cl))
        print("=" * 70, f"\n[docx] 解析出 {len(cl)} 项 → 分出 {len(chunks)} 块")
        for c in chunks:
            print(f"  #{c['order_index']} {c['type']} 章节={c['section_path'] or '—'}: "
                  f"{c['text'][:80].replace(chr(10), ' ')}")
    except ImportError:
        print("\n（未装 python-docx，跳过 Word 演示：pip install python-docx）")

    # 3) 用“模拟的 MinerU content_list”验证 MinerU 输出归一化 + 分块
    #    （含噪声页眉、相对图片路径、HTML 表格、公式、代码、列表）
    fake_mineru_output = [
        {"type": "header", "text": "公司内部资料", "page_idx": 0},                    # 噪声，应被丢
        {"type": "text", "text": "第1章 概述", "text_level": 1, "page_idx": 0},
        {"type": "text", "text": "本章介绍系统总体架构。", "page_idx": 0},
        {"type": "image", "img_path": "images/fig1.jpg", "image_caption": ["图1 架构图"], "page_idx": 0},
        {"type": "table", "img_path": "images/tab1.jpg",
         "table_body": "<html><body><table><tr><td>模块</td><td>QPS</td></tr>"
                       "<tr><td>检索</td><td>1000</td></tr></table></body></html>",
         "table_caption": ["表1 性能"], "page_idx": 1},
        {"type": "equation", "text": "$$F(x)=ax+b$$", "text_format": "latex", "page_idx": 1},
        {"type": "code", "code_body": "print('hi')", "page_idx": 1},
        {"type": "list", "list_items": ["要点一", "要点二"], "page_idx": 1},
        {"type": "footer", "text": "第 1 页", "page_idx": 1},                        # 噪声
    ]
    normalized = loader.parser.mineru_parser.normalize(fake_mineru_output, tmp / "images")
    chunks = asyncio.run(loader.asplit(normalized))
    print("=" * 70, f"\n[模拟MinerU] 原始 {len(fake_mineru_output)} 项 → 归一化 {len(normalized)} 项 "
                    f"→ 分出 {len(chunks)} 块（页眉页脚已剔除）")
    for c in chunks:
        print(f"  #{c['order_index']} {c['type']:8s} 章节={c['section_path'] or '—'}: "
              f"{c['text'][:90].replace(chr(10), ' ')}")


if __name__ == "__main__":
    _demo()
