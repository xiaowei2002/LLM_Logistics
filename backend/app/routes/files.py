"""文件上传解析接口：提取文档文字，作为模型回答的上下文。"""

from pathlib import Path

from fastapi import APIRouter, HTTPException, UploadFile

router = APIRouter(prefix="/files", tags=["files"])

ALLOWED_EXTS = {".txt", ".md", ".csv", ".log", ".pdf", ".docx", ".xlsx"}
MAX_SIZE = 10 * 1024 * 1024  # 10MB
MAX_CHARS = 30000  # 提取文字上限，避免超出模型上下文


def _read_plain(data: bytes) -> str:
    """纯文本文件解码：utf-8 失败则尝试 gbk（中文 Windows 常见）。"""
    for enc in ("utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        import io

        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import io

        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    if ext == ".xlsx":
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"## 工作表：{ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(c.strip() for c in cells):
                    lines.append(",".join(cells))
        return "\n".join(lines)

    return _read_plain(data)


@router.post("/extract")
async def extract_text(file: UploadFile) -> dict:
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型：{ext or '未知'}")

    data = await file.read()
    if len(data) > MAX_SIZE:
        raise HTTPException(status_code=400, detail="文件超过 10MB 上限")

    try:
        text = _extract(filename, data)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败：{e}")

    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="未能从文件中提取到文字（可能是扫描件/纯图片 PDF）")

    truncated = len(text) > MAX_CHARS
    return {
        "filename": filename,
        "text": text[:MAX_CHARS],
        "truncated": truncated,
        "size": len(data),
    }
